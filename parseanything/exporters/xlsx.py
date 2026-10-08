import os
import openpyxl
from openpyxl.styles import Font, PatternFill
from parseanything.schema import Document
from parseanything.interfaces import Exporter
from parseanything.registry import register_exporter

class XlsxExporter(Exporter):
    name = "xlsx"

    def export(self, doc: Document, output_path: str, options: dict) -> None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active) # remove default sheet
        
        sources_sheet = wb.create_sheet(title="Sources")
        sources_sheet.append(["Sheet Name", "Source Page", "Block ID"])
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        
        for cell in sources_sheet["1:1"]:
            cell.font = header_font
            cell.fill = header_fill
            
        has_tables = False
        table_idx = 1
        for page in doc.pages:
            for block in page.blocks:
                if block.type == "table":
                    has_tables = True
                    sheet_name = f"Table_{table_idx}"
                    ws = wb.create_sheet(title=sheet_name)
                    sources_sheet.append([sheet_name, page.number, block.id])
                    
                    if block.table and block.table.cells:
                        min_row = min(cell.row for cell in block.table.cells)
                        min_col = min(cell.col for cell in block.table.cells)
                        for cell in block.table.cells:
                            r = (cell.row - min_row) + 1
                            c = (cell.col - min_col) + 1
                            ws_cell = ws.cell(row=r, column=c, value=cell.text)
                            if getattr(cell, 'is_header', False):
                                ws_cell.font = header_font
                                ws_cell.fill = header_fill
                            if cell.rowspan > 1 or cell.colspan > 1:
                                ws.merge_cells(
                                    start_row=r, start_column=c,
                                    end_row=r + cell.rowspan - 1, end_column=c + cell.colspan - 1
                                )
                    elif block.table and block.table.markdown:
                        ws.append(["Extracted Markdown:"])
                        for row_txt in block.table.markdown.split('\n'):
                            ws.append([row_txt])
                    else:
                        ws.append(["Table content could not be structured into cells."])
                        if block.content:
                            for row_txt in block.content.split('\n'):
                                ws.append([row_txt])
                                
                    table_idx += 1
                    
        if not has_tables:
            ws = wb.create_sheet(title="Empty")
            ws.append(["No tables found in this document."])
            wb.active = ws
        else:
            # Set the first table as the active sheet instead of 'Sources'
            # so the user immediately sees their data upon opening
            wb.active = wb.worksheets[1] if len(wb.worksheets) > 1 else wb.worksheets[0]
            
        wb.save(output_path)

register_exporter(XlsxExporter())
