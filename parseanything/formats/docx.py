import os
from typing import Any
from parseanything.interfaces import FormatParser
from parseanything.schema import Document, Page, Block, BlockType, BBox, TableData, TableCell
from parseanything.registry import register_format_parser
from parseanything.errors import ErrorCode, ParseError
import docx

class DocxParser(FormatParser):
    name = "docx"
    extensions = {"docx"}
    mime = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}

    def parse(self, path: str, opts: Any) -> Document:
        doc = Document(source=os.path.basename(path), format="docx")
        page = Page(number=1, width=0, height=0, kind="virtual")
        
        try:
            docx_doc = docx.Document(path)
        except Exception as e:
            doc.errors.append(ParseError(code=ErrorCode.CORRUPT_FILE, message=str(e)))
            return doc
            
        blocks = []
        block_idx = 1
        
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        for section in docx_doc.sections:
            for p in section.header.paragraphs:
                if p.text.strip():
                    blocks.append(Block(id=f"p1_b{block_idx}", type=BlockType.HEADER, page=1, locator=f"docx:header:{block_idx}", content=p.text.strip(), reading_order=block_idx))
                    block_idx += 1
            for p in section.footer.paragraphs:
                if p.text.strip():
                    blocks.append(Block(id=f"p1_b{block_idx}", type=BlockType.FOOTER, page=1, locator=f"docx:footer:{block_idx}", content=p.text.strip(), reading_order=block_idx))
                    block_idx += 1

        for child in docx_doc.element.body:
            if isinstance(child, CT_P):
                p = Paragraph(child, docx_doc)
                if not p.text.strip():
                    continue
                level = None
                btype = BlockType.PARAGRAPH
                if p.style.name.startswith("Heading"):
                    btype = BlockType.HEADING
                    try:
                        level = int(p.style.name.split()[-1])
                    except:
                        level = 1
                elif "List" in p.style.name:
                    btype = BlockType.LIST
                
                blocks.append(Block(
                    id=f"p1_b{block_idx}",
                    type=btype,
                    page=1,
                    locator=f"docx:para:{block_idx}",
                    content=p.text,
                    level=level,
                    reading_order=block_idx
                ))
                block_idx += 1
            elif isinstance(child, CT_Tbl):
                tbl = Table(child, docx_doc)
                cells = []
                for r_idx, row in enumerate(tbl.rows):
                    for c_idx, cell in enumerate(row.cells):
                        cells.append(TableCell(
                            row=r_idx,
                            col=c_idx,
                            text=cell.text.strip()
                        ))
                
                table_data = TableData(
                    n_rows=len(tbl.rows),
                    n_cols=len(tbl.columns) if tbl.columns else 0,
                    cells=cells,
                    markdown=None
                )
                blocks.append(Block(
                    id=f"p1_b{block_idx}",
                    type=BlockType.TABLE,
                    page=1,
                    locator=f"docx:table:{block_idx}",
                    table=table_data,
                    reading_order=block_idx
                ))
                block_idx += 1

        page.blocks = blocks
        doc.pages.append(page)
        doc.stats.pages = 1
        doc.stats.blocks = len(blocks)
        return doc

register_format_parser(DocxParser())
