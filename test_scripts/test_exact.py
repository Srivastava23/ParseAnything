import json
from parseanything.schema import Document
from parseanything.formats.xlsx import XlsxParser
from parseanything.formats.csv_parser import CSVParser
from parseanything.exporters.xlsx import XlsxExporter
from parseanything import parse, Options
import pandas as pd

df = pd.DataFrame({"Name": ["A", "B"], "Age": [10, 20]})
df.to_csv("test.csv", index=False)
df.to_excel("test.xlsx", index=False)

def test_flow(file_path):
    print(f"Testing {file_path}...")
    doc = parse(file_path, options=Options())
    
    # mimic UI
    doc_json_str = json.dumps(doc.model_dump())
    
    # backend /export/xlsx
    doc_dict = json.loads(doc_json_str)
    document = Document(**doc_dict)
    
    out_path = f"export_{file_path}.xlsx"
    exporter = XlsxExporter()
    exporter.export(document, out_path, {})
    
    import openpyxl
    wb = openpyxl.load_workbook(out_path)
    print("Sheets:", wb.sheetnames)
    if len(wb.sheetnames) > 1:
        for row in wb[wb.sheetnames[1]].iter_rows(values_only=True):
            print(row)
    print("-" * 20)

test_flow("test.csv")
test_flow("test.xlsx")
