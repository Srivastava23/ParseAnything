import os
from typing import Any
from parseanything.interfaces import FormatParser
from parseanything.schema import Document, Page, Block, BlockType, TableData, TableCell
from parseanything.registry import register_format_parser
from parseanything.errors import ErrorCode, ParseError
import openpyxl

class XlsxParser(FormatParser):
    name = "xlsx"
    extensions = {"xlsx"}
    mime = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    def parse(self, path: str, opts: Any) -> Document:
        doc = Document(source=os.path.basename(path), format="xlsx")
        
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            doc.errors.append(ParseError(code=ErrorCode.CORRUPT_FILE, message=str(e)))
            return doc
            
        page_idx = 1
        block_idx = 1
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            page = Page(number=page_idx, width=0, height=0, kind="virtual")
            
            # Sheet Title
            page.blocks.append(Block(
                id=f"p{page_idx}_b{block_idx}",
                type=BlockType.HEADING,
                page=page_idx,
                content=sheet_name,
                level=1,
                reading_order=block_idx
            ))
            block_idx += 1
            
            # Table Data
            cells = []
            max_row = sheet.max_row
            max_col = sheet.max_column
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    val = sheet.cell(row=r, column=c).value
                    if val is not None:
                        cells.append(TableCell(
                            row=r-1,
                            col=c-1,
                            text=str(val)
                        ))
            
            if cells:
                table_data = TableData(
                    n_rows=max_row,
                    n_cols=max_col,
                    cells=cells
                )
                page.blocks.append(Block(
                    id=f"p{page_idx}_b{block_idx}",
                    type=BlockType.TABLE,
                    page=page_idx,
                    locator=f"xlsx:{sheet_name}",
                    table=table_data,
                    reading_order=block_idx
                ))
                block_idx += 1
            
            doc.pages.append(page)
            page_idx += 1
            
        doc.stats.pages = len(doc.pages)
        doc.stats.blocks = sum(len(p.blocks) for p in doc.pages)
        return doc

register_format_parser(XlsxParser())
