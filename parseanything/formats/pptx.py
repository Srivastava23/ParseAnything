import os
from typing import Any
from parseanything.interfaces import FormatParser
from parseanything.schema import Document, Page, Block, BlockType, BBox, TableData, TableCell
from parseanything.registry import register_format_parser
from parseanything.errors import ErrorCode, ParseError
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

class PptxParser(FormatParser):
    name = "pptx"
    extensions = {"pptx"}
    mime = {"application/vnd.openxmlformats-officedocument.presentationml.presentation"}

    def parse(self, path: str, opts: Any) -> Document:
        doc = Document(source=os.path.basename(path), format="pptx")
        
        try:
            prs = pptx.Presentation(path)
        except Exception as e:
            doc.errors.append(ParseError(code=ErrorCode.CORRUPT_FILE, message=str(e)))
            return doc
            
        block_idx = 1
        
        # pptx stores dimensions in EMU. 914400 EMU = 1 inch = 72 points
        # to convert EMU to points: EMU / 12700
        EMU_TO_PT = 12700.0
        
        for i, slide in enumerate(prs.slides):
            page_num = i + 1
            width_pt = prs.slide_width / EMU_TO_PT
            height_pt = prs.slide_height / EMU_TO_PT
            
            page = Page(number=page_num, width=width_pt, height=height_pt, kind="virtual")
            blocks = []
            
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    continue  # skip groups for simplicity or recursively handle (omitted for now)
                    
                bbox = None
                if hasattr(shape, "left") and shape.left is not None:
                    bbox = BBox(
                        x0=shape.left / EMU_TO_PT,
                        y0=shape.top / EMU_TO_PT,
                        x1=(shape.left + shape.width) / EMU_TO_PT,
                        y1=(shape.top + shape.height) / EMU_TO_PT
                    )
                
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if not text:
                        continue
                        
                    btype = BlockType.PARAGRAPH
                    if shape == slide.shapes.title:
                        btype = BlockType.HEADING
                        
                    blocks.append(Block(
                        id=f"p{page_num}_b{block_idx}",
                        type=btype,
                        page=page_num,
                        locator=f"pptx:slide{page_num}:shape{shape.shape_id}",
                        content=text,
                        bbox=bbox,
                        reading_order=block_idx
                    ))
                    block_idx += 1
                
                elif shape.has_table:
                    tbl = shape.table
                    cells = []
                    for r_idx, row in enumerate(tbl.rows):
                        for c_idx, cell in enumerate(row.cells):
                            cells.append(TableCell(
                                row=r_idx,
                                col=c_idx,
                                text=cell.text_frame.text.strip() if cell.text_frame else ""
                            ))
                    
                    table_data = TableData(
                        n_rows=len(tbl.rows),
                        n_cols=len(tbl.columns) if tbl.columns else 0,
                        cells=cells,
                        markdown=None
                    )
                    blocks.append(Block(
                        id=f"p{page_num}_b{block_idx}",
                        type=BlockType.TABLE,
                        page=page_num,
                        locator=f"pptx:slide{page_num}:table{shape.shape_id}",
                        table=table_data,
                        bbox=bbox,
                        reading_order=block_idx
                    ))
                    block_idx += 1

            page.blocks = blocks
            doc.pages.append(page)
            
        doc.stats.pages = len(prs.slides)
        doc.stats.blocks = sum(len(p.blocks) for p in doc.pages)
        return doc

register_format_parser(PptxParser())
